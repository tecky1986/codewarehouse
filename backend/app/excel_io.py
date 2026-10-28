from io import BytesIO

from openpyxl import Workbook, load_workbook


CONTAINER_TYPE_HEADERS = [
    "code",
    "name",
    "iso_code",
    "inner_l",
    "inner_w",
    "inner_h",
    "door_w",
    "door_h",
    "max_payload",
    "tare_weight",
    "volume",
    "freight_price",
    "max_stack_layers",
    "enabled",
    "remark",
]


def workbook_to_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_container_types(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "container_types"
    ws.append(CONTAINER_TYPE_HEADERS)
    for r in rows:
        ws.append([r.get(h) for h in CONTAINER_TYPE_HEADERS])
    return workbook_to_bytes(wb)


def import_container_types(content: bytes) -> list[dict]:
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for raw in rows[1:]:
        if raw is None or all(v is None or v == "" for v in raw):
            continue
        item = {}
        for i, h in enumerate(headers):
            if not h or i >= len(raw):
                continue
            item[h] = raw[i]
        if item.get("code"):
            out.append(item)
    return out


PACKAGE_MASTER_HEADERS = [
    "pn",
    "code",
    "name",
    "package_type",
    "cargo_type",
    "l",
    "w",
    "h",
    "tare",
    "net",
    "gross",
    "volume",
    "rotatable",
    "same_type_stack",
    "default_loading",
    "max_stack",
    "max_load",
    "enabled",
    "remark",
]

_PACKAGE_HEADER_ALIASES = {
    "PN": "pn",
    "PackMatls": "name",
    "packmatls": "name",
    "包装类型": "package_type",
    "货物类型": "cargo_type",
    "长cm": "_l_cm",
    "宽cm": "_w_cm",
    "高cm": "_h_cm",
    "长": "_l_cm",
    "宽": "_w_cm",
    "高": "_h_cm",
    "箱重kg": "tare",
    "净重kg": "net",
    "毛重kg": "gross",
    "四面进铲": "rotatable",
    "可旋转": "rotatable",
    "同种堆叠": "same_type_stack",
    "默认装卸": "default_loading",
    "最大堆码": "max_stack",
    "每箱承重": "max_load",
    "启用": "enabled",
    "备注": "remark",
}


def export_packages(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "packages"
    ws.append(PACKAGE_MASTER_HEADERS)
    for r in rows:
        ws.append([r.get(h) for h in PACKAGE_MASTER_HEADERS])
    return workbook_to_bytes(wb)


def package_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "packages"
    ws.append(["PN", "PackMatls", "长cm", "宽cm", "高cm", "箱重kg", "净重kg", "毛重kg"])
    ws.append(["PN-001", "OUTCASE(CTN)", 120, 80, 90, 10, 20, 30])
    return workbook_to_bytes(wb)


def _to_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip().replace(",", ""))
    except ValueError:
        return None


def _to_str(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def import_packages(content: bytes) -> list[dict]:
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    mapped = [_PACKAGE_HEADER_ALIASES.get(h, h) for h in headers]
    out = []
    for raw in rows[1:]:
        if raw is None or all(v is None or v == "" for v in raw):
            continue
        item: dict = {}
        for i, h in enumerate(mapped):
            if not h or i >= len(raw):
                continue
            item[h] = raw[i]
        pn = _to_str(item.get("pn"))
        name = _to_str(item.get("name"))
        package_type = _to_str(item.get("package_type")) or name
        l_cm = _to_num(item.get("_l_cm"))
        w_cm = _to_num(item.get("_w_cm"))
        h_cm = _to_num(item.get("_h_cm"))
        l = _to_num(item.get("l"))
        w = _to_num(item.get("w"))
        h = _to_num(item.get("h"))
        if l is None and l_cm is not None:
            l = l_cm * 10
        if w is None and w_cm is not None:
            w = w_cm * 10
        if h is None and h_cm is not None:
            h = h_cm * 10
        out.append(
            {
                "pn": pn,
                "code": _to_str(item.get("code")),
                "name": name,
                "package_type": package_type,
                "cargo_type": _to_str(item.get("cargo_type")),
                "l": int(l) if l is not None else None,
                "w": int(w) if w is not None else None,
                "h": int(h) if h is not None else None,
                "tare": _to_num(item.get("tare")),
                "net": _to_num(item.get("net")),
                "gross": _to_num(item.get("gross")),
                "volume": _to_num(item.get("volume")),
                "rotatable": item.get("rotatable"),
                "same_type_stack": item.get("same_type_stack"),
                "default_loading": _to_str(item.get("default_loading")),
                "max_stack": _to_num(item.get("max_stack")),
                "max_load": _to_num(item.get("max_load")),
                "enabled": item.get("enabled") if item.get("enabled") is not None else 1,
                "remark": _to_str(item.get("remark")),
            }
        )
    return out
