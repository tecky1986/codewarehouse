import pytest
from fastapi.testclient import TestClient

from app.deps import get_conn, reset_conn_for_tests
from app.main import app


@pytest.fixture
def client(tmp_path):
    conn = reset_conn_for_tests(tmp_path / "pkg.db")
    app.dependency_overrides[get_conn] = lambda: conn
    with TestClient(app, raise_server_exceptions=False) as c:
        yield c
    app.dependency_overrides.clear()


def test_package_crud_and_pn_unique(client):
    created = client.post(
        "/api/packages",
        json={
            "pn": "ACW1081490*80",
            "name": "OUTCASE(VEN)",
            "package_type": "OUTCASE(VEN)",
            "l": 1200,
            "w": 800,
            "h": 900,
            "tare": 40,
            "net": 200,
            "gross": 240,
            "rotatable": 1,
            "crush_sensitive": 0,
            "same_type_stack": 0,
            "default_loading": "forklift",
        },
    )
    assert created.status_code == 200
    pid = created.json()["data"]["id"]
    listed = client.get("/api/packages").json()["data"]
    assert any(p["pn"] == "ACW1081490*80" for p in listed)
    dup = client.post("/api/packages", json={"pn": "ACW1081490*80", "name": "x"})
    assert dup.status_code == 409
    updated = client.put(f"/api/packages/{pid}", json={"enabled": 0})
    assert updated.json()["data"]["enabled"] == 0
    deleted = client.delete(f"/api/packages/{pid}")
    assert deleted.status_code == 200


def test_package_excel_import_export(client):
    from io import BytesIO
    from openpyxl import Workbook, load_workbook

    tpl = client.get("/api/packages/template")
    assert tpl.status_code == 200
    assert "spreadsheetml" in tpl.headers["content-type"]

    wb = Workbook()
    ws = wb.active
    ws.append(["PN", "PackMatls", "长cm", "宽cm", "高cm", "箱重kg", "净重kg", "毛重kg"])
    ws.append(["PN-A", "CTN-A", 120, 80, 90, 10, 20, 30])
    ws.append([None, "CTN-B", 50, 40, 30, 1, 2, 3])
    ws.append(["PN-SKIP", "CTN-S", None, None, None, 1, 2, 3])
    buf = BytesIO()
    wb.save(buf)
    resp = client.post(
        "/api/packages/import",
        files={
            "file": (
                "p.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert resp.status_code == 200
    data = resp.json()["data"]
    assert data["saved_count"] == 2
    assert data["skipped"] == 1

    listed = client.get("/api/packages").json()["data"]
    by_pn = {p["pn"]: p for p in listed if p.get("pn")}
    assert by_pn["PN-A"]["l"] == 1200
    assert by_pn["PN-A"]["name"] == "CTN-A"
    no_pn = next(p for p in listed if p["name"] == "CTN-B" and not p.get("pn"))
    assert no_pn["l"] == 500

    # upsert by PN updates dimensions
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(["PN", "PackMatls", "长cm", "宽cm", "高cm", "箱重kg", "净重kg", "毛重kg"])
    ws2.append(["PN-A", "CTN-A", 150, 80, 90, 10, 20, 30])
    buf2 = BytesIO()
    wb2.save(buf2)
    again = client.post(
        "/api/packages/import",
        files={
            "file": (
                "p2.xlsx",
                buf2.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert again.status_code == 200
    listed2 = client.get("/api/packages").json()["data"]
    assert next(p for p in listed2 if p.get("pn") == "PN-A")["l"] == 1500

    exported = client.get("/api/packages/export")
    assert exported.status_code == 200
    xwb = load_workbook(BytesIO(exported.content))
    headers = [c.value for c in next(xwb.active.iter_rows(max_row=1))]
    assert "pn" in headers
    assert "name" in headers
