"""P10 round-trip: exported workbook preserves original cells except N
column (overridden) and the new 装箱顺序 column (R27/R4/R30)."""

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.deps import get_conn, reset_conn_for_tests
from app.main import app
from app.repos import ImportRowRepo

@pytest.fixture
def client(tmp_path):
    conn = reset_conn_for_tests(tmp_path / "exp.db")
    app.dependency_overrides[get_conn] = lambda: conn
    with TestClient(app, raise_server_exceptions=False) as c:
        c.conn = conn
        yield c
    app.dependency_overrides.clear()


ORIG_HEADERS = [
    "箱号 No",
    "长 Length",
    "宽 width",
    "高 high",
    "体积 CBM",
    "OUTCASE NO.",
    "PackMatls",
    "箱重",
    "净重",
    "毛重",
    "PN",
    "备注",
    "Site",
    "集装箱号",
    "关单号",
    "DN",
    "尺寸",
]


def _make_task_with_rows(client):
    body = client.post("/api/tasks", json={"name": "t"}).json()["data"]["id"]
    tid = body
    client.put(
        f"/api/tasks/{tid}",
        json={"container_mode_json": {"type_code": "20GP", "count": 2}},
    )
    conn = client.conn
    rows = [
        {
            "task_id": tid,
            "row_no": i + 1,
            "a_no": f"NO-{i+1}",
            "b_len": 1000.0,
            "c_wid": 800.0,
            "d_hgt": 600.0,
            "e_cbm": 0.48,
            "f_outcase": f"OUT-{i+1}",
            "g_packmatls": "CTN-A" if i % 2 == 0 else "CTN-B",
            "h_tare": 2.0,
            "i_net": 58.0,
            "j_gross": 60.0,
            "k_pn": f"PN-{i+1}",
            "l_note": "中文备注" if i % 2 == 0 else None,
            "m_site": None,
            "o_ref": f"REF-{i+1}",
            "p_dn": f"DN-{i+1}",
            "q_misc": None,
            "filled_fields_json": "[]",
            "source_flag": "table",
        }
        for i in range(4)
    ]
    ImportRowRepo(conn).replace_for_task(tid, rows)
    return tid


def test_export_round_trip_p10(client):
    tid = _make_task_with_rows(client)
    body = client.post(f"/api/tasks/{tid}/compute", json={"algorithms": ["dblf"]}).json()
    assert body["code"] == 0, body
    import time as _time

    deadline = _time.time() + 15
    while True:
        jbody = client.get(f"/api/jobs/{body['data']['id']}").json()
        if jbody["data"]["status"] in ("done", "failed"):
            break
        _time.sleep(0.2)
        assert _time.time() < deadline, "job did not finish"
    jbody = client.get(f"/api/jobs/{body['data']['id']}").json()
    assert jbody["data"]["status"] == "done"
    plans = client.get(f"/api/tasks/{tid}/plans").json()["data"]
    assert len(plans) == 1
    plan = plans[0]

    resp = client.get(f"/api/plans/{plan['id']}/export")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]

    wb = load_workbook(BytesIO(resp.content), data_only=True)
    ws = wb["装箱结果"]
    header = [c.value for c in ws[1]]
    # P10: original 17 headers + new column
    assert header[:17] == ORIG_HEADERS
    assert header[17] == "装箱顺序"

    conn = client.conn
    orig = conn.execute(
        "SELECT * FROM import_row WHERE task_id = ? ORDER BY row_no", (tid,)
    ).fetchall()
    data_rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    assert len(data_rows) == len(orig)

    # N column (index 13) must be a container label or empty; other cells match DB
    n_col = [r[13] for r in data_rows]
    seq_col = [r[17] for r in data_rows]
    for i, (row, dbrow) in enumerate(zip(data_rows, orig)):
        assert row[0] == dbrow["a_no"]
        # P10: B/C/D exported in cm (original table units); DB stores mm
        assert row[1] == dbrow["b_len"] / 10.0
        assert row[2] == dbrow["c_wid"] / 10.0
        assert row[3] == dbrow["d_hgt"] / 10.0
        assert row[5] == dbrow["f_outcase"]
        assert row[6] == dbrow["g_packmatls"]
        assert row[14] == dbrow["o_ref"]
        assert row[15] == dbrow["p_dn"]
        assert n_col[i] in (None, "20GP-1", "20GP-2")
        assert seq_col[i] in (None, 1, 2, 3, 4, 5, 6)

    # KPI summary sheet must exist
    assert "用箱摘要" in wb.sheetnames
    assert "未分配明细" in wb.sheetnames


def test_export_default_template_seeded(client):
    body = client.get("/api/export-templates").json()
    assert body["code"] == 0
    templates = body["data"]
    assert len(templates) >= 1
    assert any(t["is_default"] for t in templates)
    default_t = next(t for t in templates if t["is_default"])
    cols = default_t["columns_json"]
    # A-Q preserved + new column
    assert len(cols) == 18
    assert cols[-1]["field"] == "load_seq"
    # N is overridden flag
    n_col = [c for c in cols if c.get("overridden")]
    assert len(n_col) == 1


def test_default_template_cannot_be_deleted(client):
    body = client.get("/api/export-templates").json()["data"]
    default_t = next(t for t in body if t["is_default"])
    resp = client.delete(f"/api/export-templates/{default_t['id']}")
    assert resp.status_code == 409
    assert resp.json()["code"] == "E_IN_USE"


def test_export_csv_fallback_bom(client):
    tid = _make_task_with_rows(client)
    body = client.post(f"/api/tasks/{tid}/compute", json={"algorithms": ["dblf"]}).json()
    import time as _time

    deadline = _time.time() + 15
    while True:
        jbody = client.get(f"/api/jobs/{body['data']['id']}").json()
        if jbody["data"]["status"] in ("done", "failed"):
            break
        _time.sleep(0.2)
        assert _time.time() < deadline, "job did not finish"
    plans = client.get(f"/api/tasks/{tid}/plans").json()["data"]
    plan = plans[0]
    resp = client.get(f"/api/plans/{plan['id']}/export", params={"fmt": "csv"})
    assert resp.status_code == 200
    assert "text/csv" in resp.headers["content-type"]
    assert resp.content.startswith(b"\xef\xbb\xbf")


def test_chinese_filename_in_content_disposition(client):
    tid = _make_task_with_rows(client)
    body = client.post(f"/api/tasks/{tid}/compute", json={"algorithms": ["dblf"]}).json()
    import time as _time

    deadline = _time.time() + 15
    while True:
        jbody = client.get(f"/api/jobs/{body['data']['id']}").json()
        if jbody["data"]["status"] in ("done", "failed"):
            break
        _time.sleep(0.2)
        assert _time.time() < deadline, "job did not finish"
    plans = client.get(f"/api/tasks/{tid}/plans").json()["data"]
    plan = plans[0]
    resp = client.get(f"/api/plans/{plan['id']}/export")
    cd = resp.headers["content-disposition"]
    assert "filename*=UTF-8''" in cd
