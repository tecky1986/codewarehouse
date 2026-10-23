from io import BytesIO

from openpyxl import Workbook, load_workbook


CONTAINER_TYPE_HEADERS = [
    "code",
    "name",
    "iso_code",
    "inner_l",
    "inner_w",
    "inner_h",
    "door_w",
    "door_h",
    "max_payload",
    "tare_weight",
    "volume",
    "enabled",
    "remark",
]


def workbook_to_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_container_types(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "container_types"
    ws.append(CONTAINER_TYPE_HEADERS)
    for r in rows:
        ws.append([r.get(h) for h in CONTAINER_TYPE_HEADERS])
    return workbook_to_bytes(wb)


def import_container_types(content: bytes) -> list[dict]:
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for raw in rows[1:]:
        if raw is None or all(v is None or v == "" for v in raw):
            continue
        item = {}
        for i, h in enumerate(headers):
            if not h or i >= len(raw):
                continue
            item[h] = raw[i]
        if item.get("code"):
            out.append(item)
    return out
